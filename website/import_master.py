"""Importer for HYG MARKETING MASTER DATABASE.xlsx.

Understands the workbook's layout:
  - one sheet per brand (GOLDILOCKS, SAVORE, ICEBERGS, TATERS, CHATIME, ELEVATE)
  - a brand title row, then a column header row (labels located dynamically,
    since column order differs per sheet: GOLDILOCKS uses "Specific
    Instructions" while the other brands use "Date Created" there)
  - data rows mapped by the header: No. | Program / Project Name |
    Spreadsheet Link | Specific Instructions / Date Created | Created By |
    Access | Status
  - cells that carry an Excel hyperlink object (display text like
    "RAINY ADD ONS PROMO - SIT ..." over a Google Sheets URL) are resolved
    to their real URL, so links are clickable in the app
  - ALL-CAPS rows with no No. and no link are section headers; subsequent
    resources are grouped under that section (imported as a category)

Everything is normalized into the existing Brand / Category / Resource schema
-- no new models, so the rest of the system is untouched.

Run from the project root:
    python -m website.import_master "C:\\path\\to\\HYG MARKETING MASTER DATABASE.xlsx"
or via the CLI:
    flask --app app import-master --xlsx "C:\\path\\to\\file.xlsx"
"""
import datetime
import re

import openpyxl

from website.views import get_or_create_category
from website.models import Brand, Category, Resource

# Values that mean "no data" in this workbook
DROP_VALUES = {"none", "#ref!", "n/a", "-"}

# Header labels -> resource field (normalized keys)
HEADER_FIELD_MAP = {
    "no.": "ref_no",
    "no": "ref_no",
    "program / project name": "name",
    "program/project name": "name",
    "spreadsheet link": "link",
    "specific instructions": "instructions",
    "date created": "date",
    "created by": "created_by",
    "access": "access",
    "status": "status",
}

# Status values that sometimes end up in the Name column by data-entry error
STATUS_VALUES = {
    "in progress", "completed", "for review", "on hold", "not started",
    "n/a", "pending", "done", "ongoing",
}

SECTION_RE = re.compile(r"^[A-Z0-9][A-Z0-9 &,'()/.\-]*$")

DATE_PREFIX_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")


def _norm_label(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower() if value is not None else ""


def clean(value, hyperlink_target=None):
    """Normalize a cell value: strip junk, fix floats/datetimes, None when empty.

    When the cell carries an Excel hyperlink object, its target URL wins over
    the display text (e.g. "RAINY ADD ONS PROMO - SIT ..." linking to a
    Google Sheets URL).
    """
    if hyperlink_target:
        target = str(hyperlink_target).strip()
        if target.lower().startswith(("http://", "https://", "www.")):
            if target.lower().startswith("www."):
                target = "https://" + target
            return target
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.strftime("%b %d, %Y")
    if isinstance(value, datetime.date):
        return value.strftime("%b %d, %Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    text = str(value).strip()
    if not text or text.lower() in DROP_VALUES:
        return None
    return text


def _row_any(row):
    return any(v not in (None, "") for v in row)


def _name_from_url(url):
    """Readable placeholder title for rows that only contain a link."""
    if "docs.google.com/spreadsheets" in url:
        return "Google Sheet"
    if "docs.google.com/document" in url:
        return "Google Doc"
    if "docs.google.com/presentation" in url:
        return "Google Slides"
    if "docs.google.com/forms" in url:
        return "Google Form"
    if "drive.google.com" in url:
        return "Google Drive file"
    if "canva.com" in url:
        return "Canva design"
    return "(Untitled)"


def _dedupe(text):
    """Collapse accidental duplicated display text: "Coin Bank Request Coin
    Bank Request" -> "Coin Bank Request"."""
    words = text.split()
    if len(words) >= 2 and len(words) % 2 == 0:
        half = len(words) // 2
        if words[:half] == words[half:]:
            return " ".join(words[:half])
    return text


def _fallback_name(link, link_display):
    """Pick a human-friendly program name when the name cell is empty.

    Prefers the link cell's display text ("Marketing Request Data") over the
    resolved URL, so program names never show as raw Google Sheets links.
    """
    if link_display and not link_display.lower().startswith(("http://", "https://", "www.")) and link_display.lower() not in {
        "untitled spreadsheet", "untitled document", "untitled presentation",
        "untitled form", "untitled",
    }:
        return _dedupe(link_display)
    if link:
        if link.lower().startswith(("http://", "https://")):
            return _name_from_url(link)
        return link
    return "(Untitled)"


def _is_section_row(no_val, name, link):
    """ALL-CAPS title with no No. and no link -> section header."""
    if no_val or link or not name:
        return False
    return bool(SECTION_RE.match(name)) and any(ch.isalpha() for ch in name)


def _locate_header(rows):
    """Return (header_row_index, {field: col_index}) from the first header-looking row."""
    for idx, row in enumerate(rows):
        mapping = {}
        for col, value in enumerate(row):
            field = HEADER_FIELD_MAP.get(_norm_label(value))
            if field and field not in mapping:
                mapping[field] = col
        if "name" in mapping and "link" in mapping:
            return idx, mapping
    # Fallback: assume the standard layout right after the brand title row
    return 1, {"ref_no": 1, "name": 2, "link": 3, "instructions": 4,
               "date": 5, "created_by": 6, "access": 7, "status": 8}


def parse_sheet(ws, hyperlink_map=None):
    """Parse one worksheet into a list of (category_name, [resource dicts]).

    hyperlink_map maps "A1"-style cell coordinates to the URL stored in the
    cell's Excel hyperlink object (display text alone loses the URL).
    """
    hyperlink_map = hyperlink_map or {}
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    # Trim trailing fully-empty rows (CHATIME reports 1000 rows but has 8)
    while rows and not _row_any(rows[-1]):
        rows.pop()

    header_idx, colmap = _locate_header(rows)

    # Column index of the name/link fields, used for section detection on
    # sheets whose layout differs (empty name but text in the link column).
    name_col = colmap.get("name", 2)
    link_col = colmap.get("link", 3)
    no_col = colmap.get("ref_no", 1)

    def cell(row, field):
        col = colmap.get(field)
        if col is None or col >= len(row):
            return None
        coord = ws.cell(row=idx + 1, column=col + 1).coordinate
        return clean(row[col], hyperlink_map.get(coord))

    sections = []       # list of (name, [resource dicts])
    current_section = "General"
    current = []
    order = 0

    for idx, row in enumerate(rows):
        if idx <= header_idx:
            continue  # brand title row and column header row
        padded = tuple(row) + (None,) * 12

        if not _row_any(padded):
            continue

        no_val = cell(padded, "ref_no")
        name = cell(padded, "name")
        link = cell(padded, "link")
        # Display text of the link cell (before hyperlink substitution) —
        # used as a human-friendly fallback name so rows whose name cell is
        # empty don't end up titled with a raw URL.
        link_display = clean(padded[link_col]) if link_col < len(padded) else None

        # Section headers can appear in any of the three lead columns:
        #   name column:  "NOVELTIES" (GOLDILOCKS)
        #   No. column:   "COINBANK"   (GOLDILOCKS r3)
        #   link column:  "HONDA FUNRUN X ELEVATE COUPON" (GOLDILOCKS r96)
        # Rule: the cell is ALL-CAPS text, the other two lead cells are empty,
        # and it is not a URL.
        section_name = None
        lead_cells = [(no_val, no_col), (name, name_col), (link, link_col)]
        # order candidates by column so the leftmost wins when several are set
        lead_cells.sort(key=lambda item: item[1])
        filled = [(text, col) for text, col in lead_cells if text]
        if len(filled) == 1:
            text, col = filled[0]
            if "://" not in text and SECTION_RE.match(text) and any(c.isalpha() for c in text):
                # only treat as section if it sits at or left of the link column
                if col <= link_col:
                    section_name = text
        if section_name:
            if current:
                sections.append((current_section, current))
            current_section = section_name
            current = []
            continue

        instructions = cell(padded, "instructions")
        date_val = cell(padded, "date")
        created_by = cell(padded, "created_by")
        access = cell(padded, "access")
        status = cell(padded, "status")

        # Skip junk rows: no name and no link at all
        if not name and not link:
            continue
        # Skip rows where the "name" is just a date continuation
        if name and not link and DATE_PREFIX_RE.match(name):
            continue
        # Data-entry fix: a status value in the Name column -> use the link
        # cell's display text as the program name instead (e.g. GOLDILOCKS
        # row "In Progress")
        if name and name.lower() in STATUS_VALUES and link and not status:
            name, status = _fallback_name(link, link_display), name

        order += 1
        current.append({
            "ref_no": no_val,
            "name": name or _fallback_name(link, link_display),
            "link": link,
            "instructions": instructions,
            "date": date_val,
            "created_by": created_by,
            "access": access,
            "status": status,
            "sort_order": order,
        })

    if current:
        sections.append((current_section, current))
    return sections


def import_master(xlsx_path, replace=False):
    """Import every sheet of the master workbook into the database.

    replace=True wipes existing resources of the matched brands first
    (brands/categories are kept).
    """
    from website import create_app, db, BRAND_COLORS

    app = create_app()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    # Cell values alone lose the URL behind Excel hyperlink objects, so read a
    # second raw copy just to harvest hyperlink targets per sheet.
    raw_wb = openpyxl.load_workbook(xlsx_path, data_only=False)

    def _hyperlinks(ws):
        out = {}
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink is not None and c.hyperlink.target:
                    out[c.coordinate] = c.hyperlink.target
        return out

    with app.app_context():
        totals = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            links = _hyperlinks(raw_wb[sheet_name])
            brand = Brand.query.filter(
                db.func.upper(Brand.name) == sheet_name.strip().upper()
            ).first()
            if not brand:
                brand = Brand(
                    name=sheet_name.strip().upper(),
                    color=BRAND_COLORS.get(sheet_name.strip().upper(), "#333333"),
                )
                db.session.add(brand)
                db.session.flush()

            if replace:
                removed = 0
                for cat in brand.categories:
                    for res in list(cat.resources):
                        db.session.delete(res)
                        removed += 1
                db.session.flush()
                print(f"  [{brand.name}] removed {removed} existing resources")

            sections = parse_sheet(ws, links)
            brand_total = 0
            for section_name, resources in sections:
                if not resources:
                    continue
                cat = get_or_create_category(brand, section_name)
                # place new resources after existing ones
                base = len(cat.resources)
                for i, r in enumerate(resources, start=1):
                    data = {
                        "No.": r["ref_no"],
                        "Program / Project Name": r["name"],
                        "Spreadsheet Link": r["link"],
                        "Specific Instructions": r["instructions"],
                        "Date Created": r["date"],
                        "Created By": r["created_by"],
                        "Access": r["access"],
                        "Status": r["status"],
                    }
                    data = {k: v for k, v in data.items() if v is not None}
                    db.session.add(Resource(
                        ref_no=r["ref_no"],
                        name=r["name"],
                        link=r["link"],
                        instructions=r["instructions"],
                        created_by=r["created_by"],
                        access=r["access"],
                        status=r["status"],
                        sort_order=base + i,
                        category=cat,
                        data=data,
                    ))
                brand_total += len(resources)
            totals[brand.name] = brand_total

        db.session.commit()
        print("Import complete:")
        for brand_name, count in totals.items():
            print(f"  {brand_name}: {count} resources")
        return totals


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else ""
    replace_flag = "--replace" in sys.argv
    if not path:
        print("Usage: python -m website.import_master <xlsx path> [--replace]")
        raise SystemExit(1)
    import_master(path, replace=replace_flag)
